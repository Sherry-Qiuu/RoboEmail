from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "微软雅黑"
CODE_FONT = "Consolas"
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x66, 0x66, 0x66)


def set_font(run, name, size, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def para(doc, text, size=10.5, bold=False, color=INK, space_after=4, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    set_font(p.add_run(text), BODY_FONT, size, bold, color)
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(text), BODY_FONT, 12, True)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    set_font(p.add_run(text), BODY_FONT, 10.5)
    return p


def code(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.left_indent = Cm(0.4)
        set_font(p.add_run(line), CODE_FONT, 9, color=RGBColor(0x22, 0x22, 0x22))


def table(doc, rows, widths):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for index, row in enumerate(rows):
        cells = t.add_row().cells
        for col, value in enumerate(row):
            cells[col].width = Cm(widths[col])
            cp = cells[col].paragraphs[0]
            cp.paragraph_format.space_after = Pt(1)
            cp.paragraph_format.space_before = Pt(1)
            set_font(cp.add_run(value), BODY_FONT, 9.5, bold=(index == 0))
    return t


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(2)
set_font(title.add_run("RoboEmail 上线操作清单"), BODY_FONT, 19, True)

para(
    doc,
    "2026-10-29 · 全程约 40 分钟（不含域名实名审核的等待）· 详细版见 docs/deploy.md",
    9,
    color=MUTED,
    space_after=8,
)

heading(doc, "一、开始之前")
bullet(doc, "身份证、微信/支付宝、邮箱、手机，放在手边")
bullet(doc, "已经准备好的：GitHub 仓库 RoboEmail、DeepSeek key、QQ 邮箱授权码")

heading(doc, "二、先买域名（要等实名审核，所以放第一步）")
bullet(doc, "打开 dnspod.cloud.tencent.com，搜一个 .com，约 60–90 元/年")
bullet(doc, "填「域名信息模板」时，姓名和身份证号必须与实名信息完全一致")
bullet(doc, "付款后到「我的域名」看状态，等「实名认证中」变成「正常」才能解析")

heading(doc, "三、买服务器")
table(
    doc,
    [
        ("项目", "选什么"),
        ("产品", "轻量应用服务器"),
        ("地域", "中国香港（不需要备案）"),
        ("镜像", "Ubuntu 24.04 LTS，选纯净系统，不要应用镜像"),
        ("配置", "2 核 2G / 40G SSD / 20M 带宽"),
        ("时长", "1 年，约 288 元"),
        ("登录凭证", "设一个 root 密码，立刻记进记事本"),
    ],
    [3.2, 13.0],
)

heading(doc, "四、放行端口和解析域名")
bullet(doc, "控制台 → 防火墙：确认 22、80、443 三条都是「允许」")
bullet(doc, "DNS 解析加两条 A 记录：主机记录 @ 和 www，记录值填服务器公网 IP")
bullet(doc, "等 5–10 分钟，在你自己的电脑上执行下面这行，返回的地址是服务器 IP 才算生效")
code(doc, ["nslookup 你的域名"])

heading(doc, "五、连服务器、装 Docker")
code(
    doc,
    [
        "ssh root@你的服务器IP",
        "curl -fsSL https://get.docker.com | sh",
        "systemctl enable --now docker",
        "docker compose version",
    ],
)
para(doc, "最后一行要显示出 v2.x 的版本号，没有输出说明装失败了。", 9.5, color=MUTED, space_after=2)

heading(doc, "六、部署")
code(
    doc,
    [
        "mkdir -p /srv && cd /srv",
        "git clone https://github.com/Sherry-Qiuu/RoboEmail.git roboemail",
        "cd roboemail",
        "cp .env.example .env && nano .env",
    ],
)
para(
    doc,
    "在 nano 里改四项：NODE_ENV=production、APP_ORIGIN=https://你的域名、SMTP_USER / SMTP_PASS、LLM_API_KEY。"
    "存盘用 Ctrl+O 回车，退出用 Ctrl+X。",
    9.5,
    color=MUTED,
    space_after=2,
)
code(
    doc,
    [
        "nano Caddyfile      # 第一行换成你的域名，同样 Ctrl+O 回车、Ctrl+X",
        "chmod 600 .env",
        "docker compose up -d --build",
    ],
)
para(doc, "第一次要编译前端，等 3–5 分钟。结束后跑 docker compose ps，两个服务都应是 running。", 9.5, color=MUTED, space_after=2)

heading(doc, "七、验收四条")
bullet(doc, "打开 https://你的域名，是落地页，地址栏有小锁")
bullet(doc, "打开 https://你的域名/api/health，provider 显示 deepseek-chat 而不是 mock")
bullet(doc, "注册一个账号，真的收到验证码")
bullet(doc, "点忘记密码，邮件里的链接是 https://你的域名/reset?...")

heading(doc, "八、最容易翻车的三处")
table(
    doc,
    [
        ("现象", "原因", "处理"),
        ("浏览器超时，但容器是 running", "防火墙没放行 80/443", "控制台防火墙补规则"),
        ("日志里 certificate obtain failed", "DNS 还没生效", "等 10 分钟，docker compose restart caddy"),
        ("只有重置邮件链接是错的", "APP_ORIGIN 和域名不一致", "改成完全一致后重启 app"),
    ],
    [5.4, 4.6, 6.2],
)

heading(doc, "九、上线之后")
bullet(doc, "关掉本地那台机器上的临时隧道，别让开发环境长期挂在公网上")
bullet(doc, "数据库每天自动备份，先建目录再写定时任务")
code(
    doc,
    [
        "mkdir -p /srv/backup && crontab -e",
        "0 3 * * * cp /srv/roboemail/apps/server/data/roboemail.db /srv/backup/roboemail-$(date +\\%F).db",
        "0 4 * * * find /srv/backup -name 'roboemail-*.db' -mtime +14 -delete",
    ],
)
para(doc, "每天 3 点备份，保留 14 天。", 9.5, color=MUTED, space_after=4)
bullet(doc, "以后改了代码，本地 push 之后在服务器执行两行就更新完，数据库不会被覆盖")
code(doc, ["cd /srv/roboemail", "git pull && docker compose up -d --build"])

heading(doc, "十、做完发我")
para(doc, "服务器公网 IP、域名、以及卡在哪一步。密码和 key 不用发。", 10.5, space_after=2)

doc.save(r"E:\UNNC_Y1_FAM\RoboEmail\RoboEmail-上线清单.docx")
print("saved")
